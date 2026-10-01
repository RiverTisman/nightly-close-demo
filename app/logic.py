"""
Core AVRA nightly-check logic, adapted from the CLI tool to work with
uploaded files (paths to temp files) instead of fixed directory args.
No business logic changed here -- this is the same Toast parsing,
bartender pool extraction, Tip Sheet reading, and HotSchedules
cross-check as the standalone script, just callable from the web app.
"""

import csv
import json
import re
import subprocess
from collections import defaultdict, Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

from . import mapping_store
from .parse_hotschedules_roles import parse_page
from .parse_hotschedules_daily import parse_daily_roster, build_daily_roster_review, CALLIN_RE, TRAINEE_RE
from .parse_adp_punches import parse_adp_punches
from .parse_tip_sheet_daily import parse_tip_sheet_roster, read_weekly_totals_roster, TipSheetTabNotFound


# ---------------------------------------------------------------------------
# Employee mapping
# ---------------------------------------------------------------------------

class EmployeeMapping:
    def __init__(self, path):
        with open(path) as f:
            data = json.load(f)
        self.entries = data["employees"]
        self.toast_to_tipsheet = {}
        self.hotsched_to_tipsheet = {}
        self.toast_to_status = {}
        self.hotsched_to_status = {}
        self.eid_to_entry = {}
        self.hotsched_to_entry = {}
        self.tip_sheet_to_entry = {}
        self.tip_sheet_norm_to_entry = {}
        self.canonical_to_entry = {}
        for e in self.entries:
            status = e.get("status", "active")
            for alias in e.get("toast", []):
                self.toast_to_tipsheet[self._norm(alias)] = e["tip_sheet"]
                self.toast_to_status[self._norm(alias)] = status
            for alias in e.get("hotschedules", []):
                self.hotsched_to_tipsheet[self._norm(alias)] = e["tip_sheet"]
                self.hotsched_to_status[self._norm(alias)] = status
                self.hotsched_to_entry[self._norm(alias)] = e
            if e.get("eid"):
                self.eid_to_entry[self._norm_eid(e["eid"])] = e
            if e.get("tip_sheet"):
                self.tip_sheet_to_entry[e["tip_sheet"]] = e
                self.tip_sheet_norm_to_entry[self._norm(e["tip_sheet"])] = e
            if e.get("canonical"):
                self.canonical_to_entry[e["canonical"]] = e

    @staticmethod
    def _norm(name):
        return re.sub(r"\s+", " ", name.strip()).lower()

    @staticmethod
    def _norm_eid(eid):
        """Strips leading zeros so ADP's fixed-width Position ID (e.g.
        "012345") matches an EID stored from an Excel numeric cell
        ("12345")."""
        return str(eid).strip().lstrip("0") or "0"

    def toast_lookup(self, name):
        return self.toast_to_tipsheet.get(self._norm(name))

    def toast_status(self, name):
        """'active', 'terminated', or None if the name isn't in the mapping at all."""
        return self.toast_to_status.get(self._norm(name))

    def hotsched_lookup(self, name):
        return self.hotsched_to_tipsheet.get(self._norm(name))

    def hotsched_status(self, name):
        """'active', 'terminated', or None if the name isn't in the mapping at all."""
        return self.hotsched_to_status.get(self._norm(name))

    def cash_sales_number(self, tip_sheet_name):
        """The Master Server Cash Sales Sheet's "Number" for this Tip Sheet
        name, or None if that person hasn't been confidently matched to an
        entry in that sheet's Server Numbers tab yet (most of that sheet is
        old/former staff -- see [[avra-mapping-issues]])."""
        entry = self.tip_sheet_to_entry.get(tip_sheet_name)
        return entry.get("cash_sales_number") if entry else None

    def canonical_entry(self, canonical):
        """Full mapping entry (dict) for an exact canonical name, or None.

        Used by the staff-picker flow (cut list, role swaps): the client
        only ever submits a canonical it got from the picker's own staff
        list, which was built from this same mapping, so this is an exact
        lookup, not a fuzzy one -- the picker is what eliminates the
        guessing, not this method. Still returns None instead of raising
        so stale client-side data (mapping changed after the page loaded)
        surfaces as "unrecognized" rather than crashing the request."""
        return self.canonical_to_entry.get(canonical)

    def active_employees(self):
        """[{canonical, eid}, ...] for every active employee, sorted by
        canonical name -- the source list for the staff-picker UI. Kept as
        a plain list of small dicts (not full entries) since that's all
        the picker ever needs client-side, and it's what gets serialized
        into the page as JSON."""
        return sorted(
            ({"canonical": e["canonical"], "eid": e.get("eid")}
             for e in self.entries if e.get("status", "active") == "active"),
            key=lambda e: e["canonical"],
        )

    def hotsched_entry(self, name):
        """Full mapping entry (dict) for a HotSchedules alias, or None."""
        return self.hotsched_to_entry.get(self._norm(name))

    def eid_entry(self, eid):
        """Full mapping entry (dict) for an ADP EID, or None if not found.

        Normalizes leading zeros away before matching: employee_mapping.json
        stores EIDs read from Excel numeric cells (no leading zeros), but
        ADP's Position ID is a fixed-width text field that keeps them
        (e.g. "XX012345" -> "012345" once the prefix is stripped).
        """
        return self.eid_to_entry.get(self._norm_eid(eid))

    def tip_sheet_entry(self, name):
        """(entry, exact) for a literal Tip Sheet cell string, or (None,
        False). Tries an exact match first -- the real workbook's Name
        columns are exact-match Excel dropdowns (see
        build_support_staff_review), so an exact hit means this name is
        definitely the one already in employee_mapping.json. Falls back to
        a whitespace/case-normalized match so a near-miss (the sheet
        drifted slightly from our stored tip_sheet string) still resolves
        -- callers should treat `exact=False` as worth flagging, not
        silently trusting."""
        exact = self.tip_sheet_to_entry.get(name)
        if exact:
            return exact, True
        return self.tip_sheet_norm_to_entry.get(self._norm(name)), False


def _mapping_file_text(data):
    """Serializes {_comment, employees} back to the file's established
    one-employee-per-line format by hand (not json.dump, which would
    reformat every existing entry across multiple lines) -- shared by
    every function that rewrites employee_mapping.json, so a diff of any
    single change only ever shows the lines that actually changed."""
    employees = data["employees"]
    lines = ["{", f'  "_comment": {json.dumps(data.get("_comment", ""), ensure_ascii=False)},', '  "employees": [']
    for i, e in enumerate(employees):
        suffix = "," if i < len(employees) - 1 else ""
        lines.append(f"    {json.dumps(e, ensure_ascii=False)}{suffix}")
    lines.append("  ]")
    lines.append("}")
    return "\n".join(lines) + "\n"


def _save_mapping(mapping_path, data, commit_message):
    """Writes employee_mapping.json and pushes it to GitHub, returning
    the sync outcome so the page can say what really happened.

    Local write first, GitHub second: a manager whose token has expired
    still gets a working app for tonight's close, and gets told in the
    same breath that the change won't outlive the next restart. The
    reverse order would leave the running app disagreeing with the file
    it just committed."""
    text = _mapping_file_text(data)
    with open(mapping_path, "w", encoding="utf-8") as f:
        f.write(text)
    return mapping_store.push_text(text, commit_message)


def set_employee_statuses(mapping_path, updates):
    """Updates 'status' on one or more existing employees, found by exact
    canonical name -- never creates a new entry, same never-guess rule as
    everywhere else: a canonical that doesn't resolve is refused, not
    silently skipped or fuzzy-matched. updates is [{"canonical": ...,
    "status": "active"|"terminated"}, ...] (the shape the terminate/
    reactivate staff picker emits). Returns {"applied": [...], "errors":
    [...]} -- applied entries include old_status so the confirmation can
    show what actually changed."""
    with open(mapping_path, encoding="utf-8") as f:
        data = json.load(f)
    by_canonical = {e["canonical"]: e for e in data["employees"]}

    errors = []
    applied = []
    for item in updates or []:
        canonical = (item.get("canonical") or "").strip()
        status = (item.get("status") or "").strip()
        if status not in ("active", "terminated"):
            errors.append(f"{canonical or '(blank)'}: {status!r} isn't a valid status")
            continue
        entry = by_canonical.get(canonical)
        if not entry:
            errors.append(f"{canonical}: not found in employee_mapping.json")
            continue
        old_status = entry.get("status", "active")
        entry["status"] = status
        applied.append({"canonical": canonical, "old_status": old_status, "new_status": status})

    sync = None
    if applied:
        changed = ", ".join(f"{a['canonical']} -> {a['new_status']}" for a in applied)
        sync = _save_mapping(mapping_path, data, f"Set employee status: {changed}")

    return {"applied": applied, "errors": errors, "sync": sync}


def update_employee(mapping_path, original_canonical, entry):
    """Overwrites an existing employee's fields in place, found by their
    ORIGINAL canonical name (from before this edit) -- never creates a
    new entry, and refuses (rather than guesses) if that canonical
    doesn't resolve to anything.

    Exists so a manager can go back and fill in something that got
    missed the first time (a real case: an employee added without a
    HotSchedules alias, then blocked from "re-adding" it by
    add_employee_to_mapping's own duplicate-canonical refusal -- which
    was working exactly as designed, just for a goal it wasn't built to
    serve). This is that second, deliberate path: edit the one entry
    that already exists instead of trying to create a second one.

    `status` is left untouched -- that's the separate terminate/
    reactivate picker's job, not this form's."""
    with open(mapping_path, encoding="utf-8") as f:
        data = json.load(f)
    employees = data["employees"]

    original_canonical = (original_canonical or "").strip()
    target_idx = next((i for i, e in enumerate(employees) if e["canonical"] == original_canonical), None)
    if target_idx is None:
        return {"updated": False, "errors": [f"{original_canonical!r} not found in employee_mapping.json"], "sync": None}

    errors = []
    if not (entry.get("canonical") or "").strip():
        errors.append("Canonical name is required.")
    if not (entry.get("eid") or "").strip():
        errors.append("EID is required.")
    if not (entry.get("tip_sheet") or "").strip():
        errors.append("Tip Sheet name is required.")
    if errors:
        return {"updated": False, "errors": errors, "sync": None}

    entry_eid_norm = EmployeeMapping._norm_eid(entry["eid"])
    for i, e in enumerate(employees):
        if i == target_idx:
            continue
        if e.get("eid") and EmployeeMapping._norm_eid(e["eid"]) == entry_eid_norm:
            errors.append(f"EID {entry['eid']} is already used by {e['canonical']!r}")
        if e.get("tip_sheet") == entry.get("tip_sheet"):
            errors.append(f"Tip Sheet text {entry['tip_sheet']!r} is already used by {e['canonical']!r}")
        if e.get("canonical") == entry.get("canonical"):
            errors.append(f"Canonical name {entry['canonical']!r} is already used (EID {e.get('eid')})")
    if errors:
        return {"updated": False, "errors": errors, "sync": None}

    # Rebuilt in the file's established key order (matches every other
    # entry) rather than just appending "status" wherever dict insertion
    # order happens to put it -- keeps the diff to just the values that
    # actually changed instead of also reshuffling unrelated keys.
    ordered_entry = {
        "canonical": entry["canonical"],
        "tip_sheet": entry["tip_sheet"],
        "toast": entry["toast"],
        "hotschedules": entry["hotschedules"],
        "status": employees[target_idx].get("status", "active"),
        "eid": entry["eid"],
    }
    if "cash_sales_number" in entry:
        ordered_entry["cash_sales_number"] = entry["cash_sales_number"]
    employees[target_idx] = ordered_entry

    sync = _save_mapping(mapping_path, data, f"Update mapping entry for {entry['canonical']}")
    return {"updated": True, "errors": [], "canonical": entry["canonical"], "sync": sync}


def add_employee_to_mapping(mapping_path, entry):
    """Appends a new employee to employee_mapping.json and rewrites the
    file, never silently creating a duplicate -- an EID, exact tip_sheet
    string, or canonical name already in use is refused rather than added,
    the same bug class (two real employees, both
    silently-duplicated mapping entries found this project) this exists to
    prevent by construction. A duplicate canonical is included in that
    guarantee even though it's not itself a "wrong data" problem the way a
    duplicate EID/tip_sheet is -- EmployeeMapping.canonical_to_entry is a
    plain dict keyed on canonical, so two entries sharing one would make
    the second silently shadow the first everywhere canonical_entry() is
    used (the staff picker, cut confirmation, role swaps).

    Rewrites the file by hand (not json.dump, which would reformat every
    existing entry across multiple lines) to preserve the established
    one-employee-per-line convention every other entry already uses, so a
    diff of this change only ever shows the one new line."""
    with open(mapping_path, encoding="utf-8") as f:
        data = json.load(f)
    employees = data["employees"]

    errors = []
    if not (entry.get("canonical") or "").strip():
        errors.append("Canonical name is required.")
    if not (entry.get("eid") or "").strip():
        errors.append("EID is required.")
    if not (entry.get("tip_sheet") or "").strip():
        errors.append("Tip Sheet name is required.")
    if errors:
        return {"added": False, "errors": errors, "sync": None}

    entry_eid_norm = EmployeeMapping._norm_eid(entry["eid"])
    for e in employees:
        if e.get("eid") and EmployeeMapping._norm_eid(e["eid"]) == entry_eid_norm:
            errors.append(f"EID {entry['eid']} is already used by {e['canonical']!r}")
        if e.get("tip_sheet") == entry.get("tip_sheet"):
            errors.append(f"Tip Sheet text {entry['tip_sheet']!r} is already used by {e['canonical']!r}")
        if e.get("canonical") == entry.get("canonical"):
            errors.append(
                f"Canonical name {entry['canonical']!r} is already used (EID {e.get('eid')}) -- "
                "if this is genuinely a different person with the same name, add a disambiguator "
                "to the canonical name (e.g. a middle initial) so the two don't collide."
            )
    if errors:
        return {"added": False, "errors": errors, "sync": None}

    employees.append(entry)
    sync = _save_mapping(mapping_path, data, f"Add {entry['canonical']} to employee mapping")
    return {"added": True, "errors": [], "sync": sync}


def add_alias_to_employee(mapping_path, canonical, alias, system="hotschedules"):
    """Attaches one more source-system spelling to an employee who is
    already in the mapping.

    The missing third option. Adding someone is refused as a duplicate
    (correctly), and editing them means retyping their whole entry -- so
    the actual fix for "HotSchedules calls this person something we've
    never seen" used to be hand-editing the JSON. It's the single most
    common real repair: a nickname, a married name, a middle initial
    that HotSchedules started including, or a plain typo in the alias we
    stored ("DANA KOWALSKY").

    The alias string comes from a real export and is stored exactly as
    given -- no trimming beyond surrounding whitespace, no case folding.
    An alias already claimed by a DIFFERENT employee is refused rather
    than moved: two people answering to one spelling is precisely the
    ambiguity this app refuses to resolve on its own.
    """
    if system not in ("hotschedules", "toast"):
        return {"added": False, "errors": [f"{system!r} isn't a system this mapping tracks."], "sync": None}

    with open(mapping_path, encoding="utf-8") as f:
        data = json.load(f)

    canonical = (canonical or "").strip()
    alias = (alias or "").strip()
    if not alias:
        return {"added": False, "errors": ["No name to add."], "sync": None}

    target = next((e for e in data["employees"] if e["canonical"] == canonical), None)
    if target is None:
        return {"added": False,
                "errors": [f"{canonical!r} isn't in employee_mapping.json."], "sync": None}

    alias_norm = EmployeeMapping._norm(alias)
    for e in data["employees"]:
        for existing in e.get(system, []):
            if EmployeeMapping._norm(existing) != alias_norm:
                continue
            if e is target:
                return {"added": False,
                        "errors": [f"{target['canonical']} already has the {system} name "
                                   f"{existing!r}."], "sync": None}
            return {"added": False,
                    "errors": [f"{alias!r} is already the {system} name for "
                               f"{e['canonical']!r} (EID {e.get('eid')}). One spelling can only "
                               "ever point at one person -- if this really is a different person "
                               "with the same name, their exports have to be told apart first."],
                    "sync": None}

    target.setdefault(system, []).append(alias)
    sync = _save_mapping(mapping_path, data,
                         f"Add {system} name {alias!r} for {target['canonical']}")
    return {"added": True, "errors": [], "canonical": target["canonical"],
            "alias": alias, "system": system, "sync": sync}


def set_tip_sheet_name(mapping_path, canonical, tip_sheet_name):
    """Repoints an employee's `tip_sheet` at the exact string the real
    workbook holds, for when the two have drifted apart.

    Only ever called with a name read straight out of an uploaded
    workbook's WEEKLY TOTALS cell, so the result is byte-identical to
    what Excel's dropdown offers by construction. Drift is the quietest
    failure this app has -- both values look like the person's name, and
    the only symptom is a SUMIF that returns $0 -- so the fix deliberately
    doesn't route through a text box anybody could retype it into.
    """
    with open(mapping_path, encoding="utf-8") as f:
        data = json.load(f)

    canonical = (canonical or "").strip()
    target = next((e for e in data["employees"] if e["canonical"] == canonical), None)
    if target is None:
        return {"updated": False, "errors": [f"{canonical!r} isn't in employee_mapping.json."],
                "sync": None}
    if not tip_sheet_name:
        return {"updated": False, "errors": ["No Tip Sheet name given."], "sync": None}

    clash = next((e for e in data["employees"]
                  if e is not target and e.get("tip_sheet") == tip_sheet_name), None)
    if clash:
        return {"updated": False,
                "errors": [f"{tip_sheet_name!r} is already {clash['canonical']}'s Tip Sheet name "
                           f"(EID {clash.get('eid')}). Two entries pointing at one dropdown cell "
                           "would split that person's tips between them."], "sync": None}

    old = target.get("tip_sheet")
    if old == tip_sheet_name:
        return {"updated": False, "errors": ["That's already the stored Tip Sheet name."],
                "sync": None}
    target["tip_sheet"] = tip_sheet_name
    sync = _save_mapping(mapping_path, data,
                         f"Fix {target['canonical']} Tip Sheet name to match the workbook")
    return {"updated": True, "errors": [], "canonical": target["canonical"],
            "old": old, "new": tip_sheet_name, "sync": sync}


# ---------------------------------------------------------------------------
# Part 1: Toast Shift Report -> Server tip totals
# ---------------------------------------------------------------------------

EXCLUDED_TOAST_NAMES = {"bar am", "bar pm", "bar  am", "bar  pm"}


def _classify_shift_dt(dt):
    cutoff = dt.replace(hour=13, minute=0, second=0, microsecond=0)
    minutes = (dt - cutoff).total_seconds() / 60
    if abs(minutes) <= 30:
        return "AMBIGUOUS"
    return "AM" if dt < cutoff else "PM"


def classify_shift(in_date_str):
    dt = datetime.strptime(in_date_str, "%m/%d/%y %I:%M %p")
    return _classify_shift_dt(dt)


LARGE_GRATUITY_THRESHOLD = 500.0
MIN_HOURS_FOR_TIP_POOL = 4.0


def _round_half_up_int(value):
    """Cash Sales must always be a whole number for the Master Server Cash
    Sales Sheet paste -- standard rounding (.49 down, .50+ up), not
    Python's default round-half-to-even."""
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_toast_server_shifts(csv_path, mapping):
    agg = defaultdict(lambda: {"cash": 0.0, "credit": 0.0, "hours": 0.0, "toast_name": None, "large_gratuity": 0.0, "cash_sales": 0.0})
    anomalies = []
    unmatched = set()

    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name_raw = row["Employee"]
            name_norm = re.sub(r"\s+", " ", name_raw.strip()).lower()
            if name_norm in EXCLUDED_TOAST_NAMES:
                continue
            if row["Job Title"].strip().lower() != "server":
                continue

            total_sales = float(row["Total Sales"] or 0)
            hours = float(row["Hours"] or 0)
            shift = classify_shift(row["In Date"])

            if total_sales == 0 and hours > 0.1:
                anomalies.append({"name": name_raw, "reason": "zero sales but nonzero hours -- likely bad/missing clock-out",
                                   "clock_in": row["In Date"]})
                continue

            tip_sheet_name = mapping.toast_lookup(name_raw)
            if mapping.toast_status(name_raw) == "terminated":
                tip_sheet_name = f"TERMINATED EMPLOYEE APPEARS IN DATA: {name_raw}"
            elif tip_sheet_name is None:
                unmatched.add(name_raw)
                tip_sheet_name = f"UNMATCHED: {name_raw}"

            non_cash_gratuity = float(row["Non-Cash Gratuity"] or 0)
            key = (tip_sheet_name, shift)
            agg[key]["cash"] += float(row["Cash Tips Decl."] or 0)
            row_cash_sales = float(row["Cash Sales"] or 0)
            if row_cash_sales > 0:
                # Negative Cash Sales rows (refunds/voids) are never
                # subtracted from the total -- only positive sales count.
                agg[key]["cash_sales"] += row_cash_sales
            # Always added to the server's credit total automatically -- a
            # large gratuity (e.g. a private-event buyout) used to be held
            # back entirely until a manager explicitly confirmed it, which
            # made every large party a mandatory stop. Per River: default
            # to the common case (it's the server's), and surface the
            # split-with-bar-team option as a small, optional control
            # instead of a blocking review step -- see large_gratuity below
            # and the (much smaller) gratuity-resolve widget in the
            # results template.
            agg[key]["credit"] += float(row["Non-Cash Tips"] or 0) + non_cash_gratuity
            if non_cash_gratuity > LARGE_GRATUITY_THRESHOLD:
                agg[key]["large_gratuity"] += non_cash_gratuity
            agg[key]["hours"] += hours
            agg[key]["toast_name"] = name_raw

    review = []
    for (tip_sheet_name, shift), vals in sorted(agg.items(), key=lambda x: (x[0][1], x[0][0])):
        reasons = []
        if tip_sheet_name.startswith("UNMATCHED"):
            reasons.append("Unmatched Toast name -- not found in employee_mapping.json, add once confirmed")
        if tip_sheet_name.startswith("TERMINATED"):
            reasons.append("Terminated employee appears in data -- confirm before including in tip pool")
        if shift == "AMBIGUOUS":
            reasons.append("Clock-in time is near the AM/PM cutover -- shift classification is uncertain")
        # large_gratuity is informational only, not a review_reasons entry
        # -- it's already included in the credit total by default, so it
        # doesn't need needs_review/the red flag pill. The results
        # template shows it as its own small, optional "split with bar
        # team?" control instead, separate from the flagged-row mechanism.
        # PM/dinner only -- confirmed against a full week of real data that
        # AM/lunch shifts are naturally often under 4h (43% of AM shifts vs.
        # 5% of PM), so applying this to AM would exclude close to half of
        # every lunch shift for no real reason.
        excluded_from_pool = shift == "PM" and vals["hours"] < MIN_HOURS_FOR_TIP_POOL
        if excluded_from_pool:
            reasons.append(
                f"Worked {vals['hours']:.2f}h -- under the {MIN_HOURS_FOR_TIP_POOL:.0f}h tip-pool minimum. "
                "Confirm whether to add them to the tip sheet anyway."
            )
        cash_sales_number = mapping.cash_sales_number(tip_sheet_name)
        if vals["cash_sales"] > 0 and not cash_sales_number:
            reasons.append(
                "Not mapped to a Master Server Cash Sales Sheet number yet -- can't be included in "
                "the cash-sales copy buttons until employee_mapping.json has a cash_sales_number for them."
            )
        review.append({
            "shift": shift, "tip_sheet_name": tip_sheet_name, "toast_name": vals["toast_name"],
            "cash": round(vals["cash"], 2), "credit": round(vals["credit"], 2),
            "hours": round(vals["hours"], 2), "large_gratuity": round(vals["large_gratuity"], 2),
            "cash_sales": _round_half_up_int(vals["cash_sales"]), "cash_sales_number": cash_sales_number,
            "needs_review": bool(reasons), "review_reasons": reasons,
            "excluded_from_pool": excluded_from_pool,
            "exclusion_reason": "under_4h" if excluded_from_pool else None,
        })
    return review, anomalies, unmatched


def parse_toast_bartender_pool(csv_path):
    pools = {}
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name_norm = re.sub(r"\s+", " ", row["Employee"].strip()).lower()
            if name_norm not in EXCLUDED_TOAST_NAMES:
                continue
            shift = classify_shift(row["In Date"])
            pools[shift] = {
                "cash": float(row["Cash Tips Decl."] or 0),
                "credit": float(row["Non-Cash Tips"] or 0),
            }
    return pools


# ---------------------------------------------------------------------------
# Part 2: Tip Sheet role reader
# ---------------------------------------------------------------------------

TIP_SHEET_DAYS = ["MON", "TUES", "WED", "THURS", "FRI", "SAT", "SUN"]
HOTSCHED_DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
DAY_MAP = dict(zip(TIP_SHEET_DAYS, HOTSCHED_DAYS))

ROLE_SECTIONS = {
    "Server": (8, 43, 1), "Bartender": (50, 65, 1), "Runner": (72, 86, 1),
    "Barback": (72, 86, 10), "Busser": (94, 124, 1), "Barista": (94, 124, 10),
}


def load_tip_sheet_roles(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    result = {}
    for ts_day in TIP_SHEET_DAYS:
        for shift in ["AM", "PM"]:
            sheet_name = f"{ts_day} {shift}"
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for role, (r0, r1, col) in ROLE_SECTIONS.items():
                for r in range(r0, r1 + 1):
                    name = ws.cell(row=r, column=col).value
                    if not name or not str(name).strip():
                        continue
                    name = str(name).strip()
                    key = (DAY_MAP[ts_day], shift)
                    result.setdefault(name, {}).setdefault(key, set()).add(role)
    return result


# ---------------------------------------------------------------------------
# Part 3: HotSchedules PDF -> OCR -> role cross-check
# ---------------------------------------------------------------------------

def ocr_hotschedules_pdf(pdf_path, work_dir):
    """
    Rasterizes the PDF (300dpi) and runs tesseract TSV OCR on each page.
    Requires poppler-utils (pdftoppm) and tesseract-ocr to be installed
    on the host -- see Dockerfile. Returns the number of pages processed.
    """
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "300", str(pdf_path), str(work_dir / "hipage")],
        check=True, capture_output=True,
    )
    jpgs = sorted(work_dir.glob("hipage-*.jpg"))
    for jpg in jpgs:
        base = jpg.stem  # e.g. "hipage-1"
        subprocess.run(
            ["tesseract", str(jpg), str(work_dir / f"{base}_tsv"), "--psm", "6", "tsv"],
            check=True, capture_output=True,
        )
    return len(jpgs)


def parse_all_hotsched_pages(work_dir, n_pages):
    work_dir = Path(work_dir)
    all_people = []
    first_tsv = work_dir / "hipage-1_tsv.tsv"
    if not first_tsv.exists():
        return []
    _, bounds = parse_page(str(first_tsv))
    for i in range(1, n_pages + 1):
        tsv_path = work_dir / f"hipage-{i}_tsv.tsv"
        if not tsv_path.exists():
            continue
        people, _ = parse_page(str(tsv_path), bounds=bounds)
        all_people.extend(people)
    return all_people


def run_hotsched_cross_check(tip_roles, hotsched_people, mapping):
    hs_by_tipname = {}
    unmatched_hotsched_names = set()
    for p in hotsched_people:
        if not p["days"]:
            continue
        tip_name = mapping.hotsched_lookup(p["employee"])
        if tip_name is None:
            unmatched_hotsched_names.add(p["employee"])
            continue
        hs_by_tipname[tip_name] = p["days"]

    report = []
    for tip_name, day_shifts in tip_roles.items():
        hs_days = hs_by_tipname.get(tip_name)
        for (day, shift), roles in day_shifts.items():
            if tip_name not in hs_by_tipname:
                status = "NOT MATCHED / NOT FOUND IN HOTSCHEDULES"
                hs_roles = None
            else:
                shift_data = hs_days.get(day, {}).get(shift)
                if day not in hs_days:
                    status, hs_roles = "NOT SCHEDULED THAT DAY", set()
                elif shift_data is None:
                    status, hs_roles = f"NO {shift} SHIFT FOUND", set()
                else:
                    hs_roles = set(shift_data["roles"])
                    if not hs_roles:
                        status = "SHIFT FOUND, NO ROLE DETECTED"
                    elif hs_roles == roles:
                        status = "MATCH"
                    else:
                        status = "ROLE MISMATCH"
            report.append({"tip_name": tip_name, "day": day, "shift": shift,
                            "tip_roles": sorted(roles), "status": status,
                            "hs_roles": sorted(hs_roles) if hs_roles else []})
    return report, unmatched_hotsched_names


# ---------------------------------------------------------------------------
# Part 4: ADP punch report -> clock-in/out validation
# ---------------------------------------------------------------------------

SHIFT_TYPE_PATTERNS = [
    ("Opener", re.compile(r"OPEN", re.IGNORECASE)),
    ("Mid", re.compile(r"\bMID\b", re.IGNORECASE)),
    ("Closer", re.compile(r"CLOS", re.IGNORECASE)),
]


def _shift_type_from_note(note):
    """HotSchedules' per-shift note is free text a manager typed, not a
    structured field -- this only recognizes the common Opener/Mid/Closer
    convention. Returns None (never guesses) for anything else, e.g. a
    call-in note or a blank note."""
    for label, pattern in SHIFT_TYPE_PATTERNS:
        if pattern.search(note or ""):
            return label
    return None


def build_recommended_clockouts(punches, daily_roster_rows, check_date, mapping, shift_filter=None):
    """For each tipped-role ADP punch missing a clock-out on check_date,
    recommends a time based on when other people working the SAME role
    AND the same HotSchedules Opener/Mid/Closer shift type actually
    clocked out that same shift. Never invents a number when there's no
    real peer signal to average -- reports why instead (can't classify
    this person's own shift type, or no matching peer with both a
    classified shift type and a real clock-out).

    Requires the HotSchedules daily roster for shift-type classification
    -- returns [] if daily_roster_rows is empty, rather than guessing
    from ADP data alone.
    """
    if not daily_roster_rows:
        return []

    roster_by_key = {}
    for r in daily_roster_rows:
        if not r["employee"]:
            continue
        entry = mapping.hotsched_entry(r["employee"])
        if not entry or not entry.get("eid"):
            continue
        eid = EmployeeMapping._norm_eid(entry["eid"])
        roster_by_key[(eid, r["shift"])] = {
            "role": r["role"], "shift_type": _shift_type_from_note(r["note"]),
        }

    # Real, valid clock-outs that day, grouped by (role, shift, shift type)
    # so a missing-clockout person can be compared only against genuine
    # peers -- same job, same part of the day.
    valid_outs = defaultdict(list)
    for p in punches:
        if p["date"] != check_date or p["pay_code"] or p["job_role"] is None or p["out_time"] is None:
            continue
        shift = _classify_shift_dt(p["in_time"]) if p["in_time"] else None
        roster_info = roster_by_key.get((EmployeeMapping._norm_eid(p["eid"]), shift))
        shift_type = roster_info["shift_type"] if roster_info else None
        if shift_type is None:
            continue
        valid_outs[(p["job_role"], shift, shift_type)].append(p["out_time"])

    recommendations = []
    for p in punches:
        if (p["date"] != check_date or p["pay_code"] or p["job_role"] is None
                or p["out_time"] is not None or p["in_time"] is None):
            continue
        shift = _classify_shift_dt(p["in_time"])
        if shift_filter in ("AM", "PM") and shift not in (shift_filter, "AMBIGUOUS"):
            continue

        name = f"{p['first']} {p['last']}".strip()
        entry = mapping.eid_entry(p["eid"])
        roster_info = roster_by_key.get((EmployeeMapping._norm_eid(p["eid"]), shift))
        row = {
            "eid": p["eid"], "name": entry["canonical"] if entry else name,
            "job_role": p["job_role"], "shift": shift,
            "in_time": p["in_time"].strftime("%I:%M %p"),
            "recommended_out": None, "peer_count": 0,
        }

        if roster_info is None or roster_info["shift_type"] is None:
            row["shift_type"] = None
            row["note"] = (
                "Can't tell their Opener/Mid/Closer shift type from the HotSchedules note "
                "-- no recommendation possible."
            )
            recommendations.append(row)
            continue

        row["shift_type"] = roster_info["shift_type"]
        peer_outs = valid_outs.get((p["job_role"], shift, roster_info["shift_type"]), [])
        if not peer_outs:
            row["note"] = (
                f"No other {roster_info['shift_type']} {p['job_role']}s clocked out normally "
                f"that {shift} shift to compare against -- no recommendation possible."
            )
            recommendations.append(row)
            continue

        minutes = []
        for dt in peer_outs:
            m = dt.hour * 60 + dt.minute
            if dt.date() > p["in_time"].date():
                m += 24 * 60
            minutes.append(m)
        avg_minutes = round(sum(minutes) / len(minutes))
        rolled_to_next_day = avg_minutes >= 24 * 60
        avg_hour, avg_min = divmod(avg_minutes % (24 * 60), 60)
        recommended_out = f"{avg_hour % 12 or 12}:{avg_min:02d} {'AM' if avg_hour < 12 else 'PM'}"
        if rolled_to_next_day:
            recommended_out += " (next day)"

        row["recommended_out"] = recommended_out
        row["peer_count"] = len(peer_outs)
        row["note"] = (
            f"Based on {len(peer_outs)} other {roster_info['shift_type']} {p['job_role']}"
            f"{'s' if len(peer_outs) != 1 else ''} who clocked out that {shift} shift -- "
            "confirm with them before entering this, this is a suggestion, not a real punch."
        )
        recommendations.append(row)
    return recommendations


def build_clock_validation(punches, check_date, mapping, shift_filter=None):
    """Flags clock-in/out issues for one date's punches. check_date must be
    "MM/DD/YYYY" to match the ADP export's own date format. Kitchen-role
    punches (job_role is None) and non-worked pay codes (SICK/VACATION/etc)
    are skipped -- nothing to validate there.

    The under-4-hours tip-pool check is PM/dinner-only -- confirmed against
    a full week of real Toast data that AM/lunch shifts are naturally often
    under 4h (43% of AM shifts vs. 5% of PM), so applying it to AM would
    exclude close to half of every lunch shift for no real reason. Hours
    are summed per EID across all of that day's PM-classified punches (not
    just one punch row) -- confirmed ~28% of person-days have multiple
    punch rows (breaks, corrections), so a single dinner shift can show up
    as several rows that would each look "short" in isolation.
    """
    pm_hours_by_eid = defaultdict(float)
    for p in punches:
        if (p["date"] == check_date and not p["pay_code"] and p["job_role"] is not None
                and p["in_time"] and _classify_shift_dt(p["in_time"]) == "PM"):
            pm_hours_by_eid[EmployeeMapping._norm_eid(p["eid"])] += p["hours"]

    review = []
    for p in punches:
        if p["date"] != check_date or p["pay_code"] or p["job_role"] is None:
            continue

        shift = _classify_shift_dt(p["in_time"]) if p["in_time"] else None
        if shift_filter in ("AM", "PM") and shift not in (shift_filter, "AMBIGUOUS", None):
            continue

        name = f"{p['first']} {p['last']}".strip()
        entry = mapping.eid_entry(p["eid"])
        reasons = []
        if p["out_time"] is None:
            reasons.append("No clock-out recorded for this shift -- confirm it was actually closed out")
        if p["status"] == "Inactive":
            reasons.append(
                "ADP marks this punch 'Inactive' -- doesn't reliably mean terminated, but confirm before including"
            )
        if entry is None:
            reasons.append(f"EID {p['eid']} not found in employee_mapping.json -- add once confirmed")
        elif entry.get("status") == "terminated":
            reasons.append(f"TERMINATED EMPLOYEE APPEARS IN DATA: {entry['canonical']}")
        if shift == "PM":
            pm_total = pm_hours_by_eid[EmployeeMapping._norm_eid(p["eid"])]
            if pm_total < MIN_HOURS_FOR_TIP_POOL:
                reasons.append(
                    f"Worked {pm_total:.2f}h PM total today, under the {MIN_HOURS_FOR_TIP_POOL:.0f}h tip-pool "
                    "minimum -- confirm eligibility before including"
                )

        review.append({
            "eid": p["eid"], "name": name,
            "canonical": entry["canonical"] if entry else None,
            "job_code": p["job_code"], "job_role": p["job_role"], "shift": shift,
            "in_time": p["in_time"].strftime("%I:%M %p") if p["in_time"] else "",
            "out_time": p["out_time"].strftime("%I:%M %p") if p["out_time"] else "",
            "hours": round(p["hours"], 2),
            "needs_review": bool(reasons), "review_reasons": reasons,
        })
    return review


# ---------------------------------------------------------------------------
# Part 5: possible early cuts (peer clock-out comparison, all tipped roles)
# ---------------------------------------------------------------------------

# HotSchedules only records a scheduled *start* time, never an end time, so
# there's no scheduled-vs-actual comparison available. Instead this compares
# each person's clock-out against same-role, same-shift peers that night --
# confirmed against real data to be a real signal (one night had two Bussers
# clock out ~5 hours before the rest of the PM Busser group). Originally
# scoped to support staff only; broadened to every tipped role (Server,
# Bartender included) per manager request -- Host/Maitre D/Server Captain
# stay excluded since they're confirmed not Tip Sheet roles.
CUT_WATCH_ROLES = {"Server", "Bartender", "Barback", "Barista", "Busser", "Runner"}
CUT_THRESHOLD_HOURS = 2.0


def build_cut_review(punches, check_date, mapping, shift_filter=None):
    """Returns only the flagged rows -- people who clocked out
    CUT_THRESHOLD_HOURS or more before the last same-role/same-shift peer
    that night. Doesn't compute or apply a %% Cut value itself; that's a
    manager judgment call in the real Tip Sheet. Needs at least 2 people in
    a role/shift to compare against -- can't detect an outlier alone."""
    day_punches = [
        p for p in punches
        if p["date"] == check_date and p["job_role"] in CUT_WATCH_ROLES
        and p["out_time"] and not p["pay_code"]
    ]

    groups = defaultdict(list)
    for p in day_punches:
        shift = _classify_shift_dt(p["in_time"]) if p["in_time"] else None
        groups[(p["job_role"], shift)].append(p)

    review = []
    for (role, shift), members in groups.items():
        if shift_filter in ("AM", "PM") and shift not in (shift_filter, "AMBIGUOUS", None):
            continue
        if len(members) < 2:
            continue
        latest = max(m["out_time"] for m in members)
        for p in members:
            gap_hours = (latest - p["out_time"]).total_seconds() / 3600
            if gap_hours < CUT_THRESHOLD_HOURS:
                continue
            entry = mapping.eid_entry(p["eid"])
            name = f"{p['first']} {p['last']}".strip()
            review.append({
                "eid": p["eid"], "name": entry["canonical"] if entry else name,
                "job_role": role, "shift": shift,
                "out_time": p["out_time"].strftime("%I:%M %p"),
                "latest_peer_out": latest.strftime("%I:%M %p"),
                "gap_hours": round(gap_hours, 1),
                "needs_review": True,
                "review_reasons": [
                    f"Clocked out {gap_hours:.1f}h before the last {role} that shift "
                    f"({latest.strftime('%I:%M %p')}) -- possibly cut early, consider a % Cut in the Tip Sheet"
                ],
            })
    return review


# ---------------------------------------------------------------------------
# Part 6: HotSchedules-scheduled role vs. ADP job-code cross-check
# ---------------------------------------------------------------------------

# Only tip-sheet roles get cross-checked here. Host/Maitre D/Server Captain
# aren't on the Tip Sheet at all (confirmed with manager) -- they're only
# checked for clock-in/out correctness elsewhere (build_clock_validation),
# not role placement.
JOB_CODE_CROSS_CHECK_SKIP_ROLES = {"Host"}


def build_job_code_review(daily_roster_rows, punches, check_date, mapping, shift_filter=None):
    """Cross-checks each HotSchedules-scheduled tipped-role person against
    what job code they actually clocked in under in ADP that night."""
    # Keyed by EID only -- see build_support_staff_review for why matching
    # by (eid, shift) is wrong (scheduled-shift label vs. clock-time-derived
    # shift can legitimately disagree and cause a missed match).
    adp_by_eid = defaultdict(list)
    for p in punches:
        if p["date"] != check_date or p["pay_code"]:
            continue
        adp_by_eid[EmployeeMapping._norm_eid(p["eid"])].append(p)

    review = []
    for r in daily_roster_rows:
        role = r["role"]
        if role in JOB_CODE_CROSS_CHECK_SKIP_ROLES or not r["employee"]:
            continue
        if shift_filter in ("AM", "PM") and r["shift"] not in (shift_filter, "AMBIGUOUS"):
            continue
        entry = mapping.hotsched_entry(r["employee"])
        if entry is None or not entry.get("eid"):
            continue  # unmatched name / no EID yet -- already flagged elsewhere, can't cross-check

        eid = EmployeeMapping._norm_eid(entry["eid"])
        matches = adp_by_eid.get(eid, [])
        if not matches:
            continue  # scheduled but no punch that shift -- a no-show/cut question, not a job-code one
        adp_roles = sorted({p["job_role"] for p in matches if p["job_role"]})
        if adp_roles and role not in adp_roles:
            review.append({
                "name": entry["canonical"], "shift": r["shift"],
                "scheduled_role": role, "adp_roles": adp_roles,
                "needs_review": True,
                "review_reasons": [
                    f"Scheduled as {role} in HotSchedules but clocked in under "
                    f"{'/'.join(adp_roles)} in ADP -- confirm correct placement in the Tip Sheet"
                ],
            })
    return review


# ---------------------------------------------------------------------------
# Part 6.5: support-staff roster (Busser/Runner/Barista/Barback) for the
# pastable Tip Sheet, with on-call/no-show confirmation against ADP
# ---------------------------------------------------------------------------

# These four roles have no Toast record at all (unlike Server/Bartender) --
# HotSchedules (who's scheduled) and ADP (who actually punched in) are the
# only two signals available for "did this person work." The real Tip
# Sheet's Name column for these sections is what this populates; the %%
# Cut column stays the sheet's own default/manual entry, same as
# build_cut_review elsewhere -- this doesn't invent a cut percentage.
SUPPORT_STAFF_ROLES = {"Busser", "Runner", "Barista", "Barback"}


def build_support_staff_review(daily_roster_rows, punches, check_date, mapping,
                                clock_review, cut_review, shift_filter=None, role_overrides=None):
    """Builds the paste-ready roster for the four support-staff Tip Sheet
    sections. Only needs the HotSchedules daily roster -- ADP is optional,
    same as it already is for the rest of Daily Close. Pass punches=[],
    clock_review=[], cut_review=[], check_date=None when ADP wasn't
    uploaded; every row's `has_adp_punch` comes back None ("not checked")
    rather than a false "No", so the results page can tell "confirmed no
    punch" apart from "never had ADP to check against."

    A HotSchedules note containing "call in" (or "called in") means the
    shift wasn't on the original schedule -- someone was called in ad hoc,
    with no guarantee they actually showed. A [House Shift] with a name
    attached is the same uncertainty from the other direction -- someone
    picked up an open shift, not confirmed they worked it. A note
    containing "training"/"not live"/"test" is the same uncertainty from
    a third direction -- a trainee who may or may not have actually
    worked (or been paid for) the shift. Per River: none of these three
    kinds of row may ever be silently included OR excluded from the
    pastable roster -- all three default to excluded and need an explicit
    manager confirmation (handled client-side on the results page, since
    this app keeps no state between requests); ADP presence, when
    available, is only ever a double-check on whatever the manager
    selects, never an override -- a manager can confirm "worked" with no
    ADP punch, or "didn't work" despite one, and both just surface a
    warning rather than blocking the choice.

    A terminated employee is the one exception to "never silently include
    or exclude" above -- their row is excluded automatically, no
    confirmation prompt, since "did a fired person work this shift" isn't
    a real question. Still flagged for a manager's attention in case it's
    actually a data problem (rehire, stale HotSchedules profile).

    Rows are sorted by (role, shift, employee name) -- alphabetical within
    each role+shift group, matching how servers are already sorted.

    role_overrides is an optional {canonical: new_role} dict (from
    apply_role_swaps) for anyone who actually worked a different
    support-staff role than HotSchedules scheduled them for -- the
    pastable roster reflects the actual role worked, not the scheduled
    one, and the row moves to that role's section.
    """
    role_overrides = role_overrides or {}
    adp_checked = check_date is not None
    # Keyed by EID only, not (EID, shift) -- HotSchedules' shift label is
    # what was *scheduled*, while a punch's shift here would be *computed
    # from actual clock-in time* (before/after 1pm). Those two disagree
    # any time someone's real clock-in crosses that boundary (e.g.
    # scheduled AM but didn't punch in until 3:28pm), which produced a
    # false "no ADP punch" for someone who genuinely worked and punched
    # in -- confirmed against real data). The question this answers is just "did they
    # punch in that date," so shift shouldn't gate it at all.
    adp_by_eid = defaultdict(list)
    for p in punches:
        if p["date"] != check_date or p["pay_code"]:
            continue
        adp_by_eid[EmployeeMapping._norm_eid(p["eid"])].append(p)

    clock_flags_by_key = defaultdict(list)
    for r in clock_review:
        if r["eid"]:
            clock_flags_by_key[(EmployeeMapping._norm_eid(r["eid"]), r["shift"])].extend(r["review_reasons"])
    cut_flags_by_key = defaultdict(list)
    for r in cut_review:
        if r["eid"]:
            cut_flags_by_key[(EmployeeMapping._norm_eid(r["eid"]), r["shift"])].extend(r["review_reasons"])

    review = []
    for r in daily_roster_rows:
        name = r["employee"]
        if not name:
            continue  # unclaimed house shift, nobody to review
        entry = mapping.hotsched_entry(name)

        role = r["role"]
        role_overridden = False
        if entry and entry["canonical"] in role_overrides:
            role = role_overrides[entry["canonical"]]
            role_overridden = True

        if role not in SUPPORT_STAFF_ROLES:
            continue  # not a support-staff row (scheduled or overridden)
        if shift_filter in ("AM", "PM") and r["shift"] not in (shift_filter, "AMBIGUOUS"):
            continue

        eid = EmployeeMapping._norm_eid(entry["eid"]) if entry and entry.get("eid") else None
        has_adp_punch = bool(eid and adp_by_eid.get(eid)) if adp_checked else None
        is_on_call = bool(CALLIN_RE.search(r["note"] or ""))
        # A House Shift with a name attached means someone picked it up --
        # but HotSchedules doesn't confirm they actually showed, same
        # uncertainty as a call-in. Defaults to excluded from the pastable
        # roster until a manager confirms, same UX as on-call rows.
        is_house_shift = r.get("flag") == "[House Shift]"
        # A trainee note ("TRAINING", "NOT LIVE", "TEST") means HotSchedules
        # doesn't confirm they actually worked the shift either -- same
        # uncertainty as a call-in or a House Shift, so it gets the same
        # default-excluded, confirm-worked/didn't-work treatment rather
        # than silently being included or dropped.
        is_trainee = bool(TRAINEE_RE.search(r["note"] or ""))
        # A terminated employee showing up in a HotSchedules export (call-in
        # or House Shift or otherwise) isn't a "did they work" question a
        # manager needs to answer -- they're not employed here anymore, so
        # it's excluded automatically rather than prompting for confirmation
        # the same way an active on-call/House Shift row would. Still
        # flagged, since it's worth a manager's eyes as a possible data
        # error (stale HotSchedules profile, or an actual rehire).
        is_terminated = entry is not None and entry.get("status") == "terminated"

        adp_flags = []
        if eid:
            adp_flags += clock_flags_by_key.get((eid, r["shift"]), [])
            adp_flags += cut_flags_by_key.get((eid, r["shift"]), [])
        if entry is None:
            adp_flags.append(f"Unmatched name: {name!r} -- not found in employee_mapping.json")
        if is_terminated:
            adp_flags.append(
                f"TERMINATED EMPLOYEE APPEARS IN DATA: {entry['canonical']} -- excluded from the pastable "
                "roster automatically. Confirm this isn't a rehire (flip their status back) or a scheduling error."
            )
        if role_overridden:
            adp_flags.append(f"Manager reported this as a {role} shift tonight (scheduled as {r['role']})")

        # Pasted/displayed name must be the exact Tip Sheet string (same
        # rule as servers) -- the real workbook's Name columns are Excel
        # data-validation dropdowns sourced from WEEKLY TOTALS, which only
        # register a paste as "selected" on an exact byte-for-byte match
        # (confirmed real quirk: several real names have a trailing space
        # or other stray character in that source list). `canonical` is a
        # cleaned-up display name for OUR purposes, not what the sheet
        # expects -- pasting it silently breaks the dropdown match.
        tip_sheet_name = entry.get("tip_sheet") if entry else None
        review.append({
            "shift": r["shift"], "role": role,
            "employee": tip_sheet_name or f"UNMATCHED: {name}",
            "eid": eid, "has_adp_punch": has_adp_punch, "is_on_call": is_on_call,
            "is_house_shift": is_house_shift, "is_trainee": is_trainee, "role_overridden": role_overridden,
            "is_terminated": is_terminated,
            "hotsched_note": r["note"],
            "adp_flags": list(dict.fromkeys(adp_flags)),
            "needs_review": is_on_call or is_house_shift or is_trainee or role_overridden or is_terminated or bool(adp_flags) or not tip_sheet_name,
        })
    review.sort(key=lambda r: (r["role"], r["shift"], r["employee"]))
    return review


# ---------------------------------------------------------------------------
# Part 6.75: previous day's real Tip Sheet vs. ADP -- daily payroll audit
# ---------------------------------------------------------------------------

# Only these six roles are ever on the Tip Sheet at all (Host/Maitre D/
# Server Captain are confirmed clock-only elsewhere, see
# JOB_CODE_CROSS_CHECK_SKIP_ROLES) -- this is also the display order,
# matching the real sheet's own section order.
TIP_SHEET_ROLE_ORDER = ["Server", "Bartender", "Runner", "Barback", "Busser", "Barista"]


def build_tip_sheet_reconciliation(combined_roster, punches, check_date, mapping, shift_filter=None):
    """The daily audit River already does by hand: print the Punch Source
    Report and yesterday's Tip Sheet, and check each against the other.
    combined_roster is {(role, shift): [tip sheet name strings]} -- one
    entry per day+shift tab actually uploaded (see parse_tip_sheet_roster).

    Three things this catches, matching River's own description of the
    process:
    - someone's on the Tip Sheet with no matching ADP punch that shift --
      they may not have actually worked, remove them or they get paid for
      a day they didn't work
    - someone punched in under a tipped role in ADP but isn't on the Tip
      Sheet anywhere -- add them, or they don't get paid at all
    - someone's Tip Sheet section doesn't match what ADP shows they
      clocked in under -- wrong department, pays incorrectly

    Also folds in the same missing-clock-out/Inactive/terminated signals
    as build_clock_validation, so this one list covers "all clock ins are
    correct" too, not just presence/absence.
    """
    # adp_by_eid_shift still buckets by shift -- needed below to list each
    # unclaimed punch under the shift it actually happened. adp_by_eid
    # (eid only) is for the "does this Tip Sheet name have a punch at all"
    # check just below: matching by (eid, shift) there was wrong, since
    # the Tip Sheet's shift is which tab/day it's printed on while a
    # punch's shift here is derived from actual clock-in time, and those
    # two can legitimately disagree (see build_support_staff_review).
    adp_by_eid_shift = defaultdict(list)
    adp_by_eid = defaultdict(list)
    for p in punches:
        if p["date"] != check_date or p["pay_code"] or p["job_role"] is None:
            continue
        shift = _classify_shift_dt(p["in_time"]) if p["in_time"] else None
        adp_by_eid_shift[(EmployeeMapping._norm_eid(p["eid"]), shift)].append(p)
        adp_by_eid[EmployeeMapping._norm_eid(p["eid"])].append(p)

    rows = []
    seen_eid_shift = set()

    for (role, shift), names in combined_roster.items():
        if shift_filter in ("AM", "PM") and shift != shift_filter:
            continue
        for name in names:
            entry, exact = mapping.tip_sheet_entry(name)
            reasons = []
            eid = EmployeeMapping._norm_eid(entry["eid"]) if entry and entry.get("eid") else None

            if entry is None:
                reasons.append(
                    f"Tip Sheet name {name!r} not found in employee_mapping.json -- "
                    "can't cross-check, add it once confirmed"
                )
            elif not exact:
                reasons.append(
                    f"Tip Sheet has {name!r}, mapping has {entry['tip_sheet']!r} -- "
                    "close but not an exact match, may need re-syncing"
                )

            matches = adp_by_eid.get(eid, []) if eid else []
            if eid and not matches:
                reasons.append(
                    f"On the Tip Sheet under {role} but no ADP punch found for {shift} on "
                    f"{check_date} -- confirm they actually worked, remove if not"
                )
            elif matches:
                adp_roles = sorted({p["job_role"] for p in matches if p["job_role"]})
                if adp_roles and role not in adp_roles:
                    reasons.append(
                        f"On the Tip Sheet under {role} but ADP shows "
                        f"{'/'.join(adp_roles)} -- wrong department, will be paid incorrectly"
                    )
                for p in matches:
                    if p["out_time"] is None:
                        reasons.append("No clock-out recorded for this shift -- confirm it was actually closed out")
                    if p["status"] == "Inactive":
                        reasons.append("ADP marks this punch 'Inactive' -- confirm before including")
                if entry and entry.get("status") == "terminated":
                    reasons.append(f"TERMINATED EMPLOYEE APPEARS IN DATA: {entry['canonical']}")

            if eid:
                # Mark both the Tip Sheet's own shift and every shift they
                # actually punched under as seen -- otherwise a person whose
                # real clock-in landed in a different shift bucket than the
                # Tip Sheet tab would also show up below as an "unclaimed"
                # ADP punch, duplicating them under the other shift.
                seen_eid_shift.add((eid, shift))
                for p in matches:
                    seen_eid_shift.add((eid, _classify_shift_dt(p["in_time"]) if p["in_time"] else None))

            rows.append({
                "role": role, "shift": shift,
                "name": entry["canonical"] if entry else name,
                "on_tip_sheet": True, "has_adp_punch": bool(matches),
                "needs_review": bool(reasons),
                "review_reasons": list(dict.fromkeys(reasons)),
            })

    for (eid, shift), matches in adp_by_eid_shift.items():
        if shift_filter in ("AM", "PM") and shift != shift_filter:
            continue
        if (eid, shift) in seen_eid_shift:
            continue
        role = matches[0]["job_role"]
        if role not in TIP_SHEET_ROLE_ORDER:
            continue  # kitchen or Host/Maitre D/Server Captain -- not a Tip Sheet role
        entry = mapping.eid_entry(eid)
        p = matches[0]
        rows.append({
            "role": role, "shift": shift,
            "name": entry["canonical"] if entry else f"{p['first']} {p['last']}".strip(),
            "on_tip_sheet": False, "has_adp_punch": True,
            "needs_review": True,
            "review_reasons": [
                f"Clocked in under {role} in ADP on {check_date} but not found anywhere "
                "on the Tip Sheet -- add them or they won't get paid"
            ],
        })

    rows.sort(key=lambda r: (
        TIP_SHEET_ROLE_ORDER.index(r["role"]) if r["role"] in TIP_SHEET_ROLE_ORDER else 99,
        r["shift"], r["name"],
    ))
    return rows


def build_mapping_dropdown_check(weekly_totals_roster, mapping):
    """Cross-checks every active, EID-linked employee_mapping.json entry
    against a real Tip Sheet workbook's WEEKLY TOTALS roster -- the exact
    source Excel's own Name-column dropdowns pull from. Hard rule (River,
    2026-07-23): tip_sheet must byte-match that dropdown text always, no
    exceptions, even when the sheet's own text is factually wrong -- a
    mismatch doesn't just fail to paste, it makes the sheet's own
    SUMIF-based totals silently return $0 for that person's entire
    tip-out share on every role their tips flow through. Confirmed real
    case: one Runner's Runner-role mismatch silently
    dropped his entire $10.64 cash + $189.24 credit share from the
    7/22/26 close's grand total.

    Never auto-fixes anything -- just flags exactly what's wrong (which
    canonical name, what's mapped vs. what the real sheet says) so it can
    be corrected in employee_mapping.json.
    """
    real_by_eid = {}
    for eid, name in weekly_totals_roster:
        real_by_eid[EmployeeMapping._norm_eid(eid)] = name

    mismatches = []
    not_found = []
    checked = 0
    for e in mapping.entries:
        if e.get("status") != "active" or not e.get("eid"):
            continue
        checked += 1
        eid_norm = EmployeeMapping._norm_eid(e["eid"])
        real_name = real_by_eid.get(eid_norm)
        if real_name is None:
            not_found.append({"canonical": e["canonical"], "eid": e["eid"]})
            continue
        if e.get("tip_sheet") != real_name:
            mismatches.append({
                "canonical": e["canonical"], "eid": e["eid"],
                "mapped": e.get("tip_sheet"), "real": real_name,
            })
    return {"checked_count": checked, "mismatches": mismatches, "not_found": not_found}


# ---------------------------------------------------------------------------
# Part 7: manager-reported cuts (primary source) vs. backend signals (check)
# ---------------------------------------------------------------------------
#
# Both the cut list and role swaps come from the staff-picker UI (see
# static/staff_picker.js) rather than free text: the manager searches and
# selects from the real active-employee list, so what the server receives
# is always an exact canonical name already in employee_mapping.json --
# never text to fuzzy-match. This is what makes the "names never match"
# bug class structurally impossible here, rather than merely less likely.
# The one remaining failure mode -- a canonical the client sent that no
# longer exists server-side (mapping edited between page load and submit)
# -- still can't be silently dropped, so it's surfaced as unrecognized
# same as before.

def apply_role_swaps(role_swaps, mapping):
    """role_swaps is [{"canonical": ..., "role": ...}, ...] from the
    picker. Returns ({canonical: new_role}, unrecognized) where
    unrecognized holds any entry whose canonical or role no longer
    resolves -- stale client-side data, not a name-matching failure."""
    overrides = {}
    unrecognized = []
    for item in role_swaps or []:
        canonical = (item.get("canonical") or "").strip()
        role = (item.get("role") or "").strip()
        if not mapping.canonical_entry(canonical) or role not in SUPPORT_STAFF_ROLES:
            unrecognized.append(item)
            continue
        overrides[canonical] = role
    return overrides, unrecognized


def build_cut_confirmation(cut_canonicals, cut_review, server_review, clock_review, mapping):
    """The manager's reported cut list is the primary source of truth --
    this cross-references it against the backend's automated signals
    (peer-comparison cut_review, under-4h PM exclusions) purely as a
    double-check, never to override what the manager reported."""
    manager_canonicals = set()
    unrecognized = []
    for canonical in cut_canonicals or []:
        canonical = (canonical or "").strip()
        if mapping.canonical_entry(canonical):
            manager_canonicals.add(canonical)
        else:
            unrecognized.append(canonical)

    backend_candidates = defaultdict(list)
    for r in cut_review:
        backend_candidates[r["name"]].append(
            f"Clocked out {r['gap_hours']}h before the last {r['job_role']} that shift"
        )
    for r in server_review:
        if r.get("excluded_from_pool"):
            backend_candidates[r["tip_sheet_name"]].append(
                f"Toast: worked only {r['hours']:.2f}h (PM)"
            )
    for r in clock_review:
        if any("tip-pool minimum" in x for x in r["review_reasons"]):
            backend_candidates[r["canonical"] or r["name"]].append("ADP: under 4h PM total")

    for name in backend_candidates:
        backend_candidates[name] = list(dict.fromkeys(backend_candidates[name]))

    confirmed = [
        {"name": name, "reasons": backend_candidates[name]}
        for name in sorted(manager_canonicals) if name in backend_candidates
    ]
    manager_only = [
        {"name": name} for name in sorted(manager_canonicals) if name not in backend_candidates
    ]
    backend_only = [
        {"name": name, "reasons": reasons} for name, reasons in sorted(backend_candidates.items())
        if name not in manager_canonicals
    ]

    return {
        "manager_reported": bool(cut_canonicals),
        "excluded_canonicals": manager_canonicals,
        "unrecognized": unrecognized,
        "confirmed": confirmed,
        "manager_only": manager_only,
        "backend_only": backend_only,
    }


# ---------------------------------------------------------------------------
# Top-level orchestration for the web app
# ---------------------------------------------------------------------------

def run_full_check(toast_csv_path, tip_sheet_path, hotsched_pdf_path, mapping_path, work_dir,
                    hotsched_daily_path=None, roster_date=None, shift_filter=None,
                    adp_punches_path=None, cut_canonicals=None, tip_sheet_daily_path=None,
                    role_swaps=None):
    """Returns a single dict with everything the results page needs.

    shift_filter is "AM", "PM", or None/"BOTH" for the full day. When
    filtering to one shift, rows the app couldn't confidently classify
    (AMBIGUOUS clock-in times) are always kept rather than risk hiding
    something that needs review -- e.g. running an AM-only lunch close
    still shows an AMBIGUOUS row even though it isn't a clean "AM" row.
    """
    mapping = EmployeeMapping(mapping_path)
    result = {}
    result["shift_filter"] = shift_filter if shift_filter in ("AM", "PM") else "BOTH"

    if toast_csv_path:
        review, anomalies, unmatched_toast = parse_toast_server_shifts(toast_csv_path, mapping)
        pools = parse_toast_bartender_pool(toast_csv_path)
        if result["shift_filter"] in ("AM", "PM"):
            review = [r for r in review if r["shift"] in (result["shift_filter"], "AMBIGUOUS")]
            pools = {s: p for s, p in pools.items() if s in (result["shift_filter"], "AMBIGUOUS")}
        result["server_review"] = review
        result["toast_anomalies"] = anomalies
        result["unmatched_toast_names"] = sorted(unmatched_toast)
        result["bartender_pools"] = pools

    if hotsched_daily_path:
        rows = parse_daily_roster(hotsched_daily_path)
        daily_review = build_daily_roster_review(rows, mapping)
        if result["shift_filter"] in ("AM", "PM"):
            daily_review = [r for r in daily_review if r["shift"] == result["shift_filter"]]
        result["daily_roster_review"] = daily_review
        result["daily_roster_date"] = roster_date
        result["daily_roster_needs_review_count"] = sum(1 for e in daily_review if e["needs_review"])

    punches = []
    adp_date = None
    date_obj = None
    if roster_date:
        try:
            date_obj = datetime.strptime(roster_date, "%Y-%m-%d")
        except ValueError:
            date_obj = None

    if adp_punches_path:
        punches = parse_adp_punches(adp_punches_path)
        if date_obj:
            adp_date = date_obj.strftime("%m/%d/%Y")
        if adp_date:
            clock_review = build_clock_validation(punches, adp_date, mapping, result["shift_filter"])
            result["clock_review"] = clock_review
            result["clock_review_date"] = adp_date
            result["clock_review_needs_review_count"] = sum(1 for e in clock_review if e["needs_review"])
            result["cut_review"] = build_cut_review(punches, adp_date, mapping, result["shift_filter"])
            if result.get("daily_roster_review") is not None:
                result["job_code_review"] = build_job_code_review(
                    result["daily_roster_review"], punches, adp_date, mapping, result["shift_filter"]
                )
                result["recommended_clockouts"] = build_recommended_clockouts(
                    punches, result["daily_roster_review"], adp_date, mapping, result["shift_filter"]
                )
            elif any(not r["out_time"] for r in clock_review):
                # Missing clock-outs exist but there's no HotSchedules daily
                # roster to classify people's Opener/Mid/Closer shift type
                # against -- can't build a real recommendation without it,
                # so ask for it explicitly rather than silently doing nothing.
                result["missing_clockouts_need_hotsched"] = True
            if tip_sheet_daily_path:
                shifts_to_check = ["AM", "PM"] if result["shift_filter"] == "BOTH" else [result["shift_filter"]]
                combined_roster = {}
                tab_errors = []
                for shift in shifts_to_check:
                    try:
                        roster = parse_tip_sheet_roster(tip_sheet_daily_path, date_obj, shift)
                        for role, names in roster.items():
                            combined_roster[(role, shift)] = names
                    except TipSheetTabNotFound as e:
                        tab_errors.append(str(e))
                if combined_roster:
                    result["tip_sheet_reconciliation"] = build_tip_sheet_reconciliation(
                        combined_roster, punches, adp_date, mapping, result["shift_filter"]
                    )
                    result["tip_sheet_reconciliation_needs_review_count"] = sum(
                        1 for r in result["tip_sheet_reconciliation"] if r["needs_review"]
                    )
                if tab_errors:
                    result["tip_sheet_reconciliation_errors"] = tab_errors
        else:
            result["clock_review_missing_date"] = True

    if tip_sheet_daily_path:
        # Independent of ADP/date -- WEEKLY TOTALS carries every currently
        # active employee regardless of which week's workbook this is, so
        # this runs any time a real Tip Sheet gets uploaded here at all.
        try:
            weekly_totals_roster = read_weekly_totals_roster(tip_sheet_daily_path)
            result["mapping_dropdown_check"] = build_mapping_dropdown_check(weekly_totals_roster, mapping)
        except TipSheetTabNotFound as e:
            result["mapping_dropdown_check_error"] = str(e)

    role_overrides, role_swap_unrecognized = apply_role_swaps(role_swaps, mapping)
    result["role_swap_unrecognized"] = role_swap_unrecognized
    result["role_swap_applied"] = [
        {"canonical": canonical, "new_role": new_role} for canonical, new_role in sorted(role_overrides.items())
    ]

    if result.get("daily_roster_review") is not None:
        # ADP is optional here -- ambient clock_review/cut_review, empty
        # otherwise. adp_date is only non-None when ADP was uploaded AND
        # matched to a valid date, so this naturally covers "no ADP" and
        # "ADP uploaded but date missing/invalid" the same way.
        result["support_staff_review"] = build_support_staff_review(
            result["daily_roster_review"], punches, adp_date, mapping,
            result.get("clock_review", []), result.get("cut_review", []), result["shift_filter"],
            role_overrides=role_overrides,
        )
        result["support_staff_adp_checked"] = adp_date is not None

    if cut_canonicals:
        cc = build_cut_confirmation(
            cut_canonicals,
            result.get("cut_review", []),
            result.get("server_review", []),
            result.get("clock_review", []),
            mapping,
        )
        result["cut_confirmation"] = cc
        if result.get("server_review") is not None:
            for r in result["server_review"]:
                entry = None
                for e in mapping.entries:
                    if e.get("tip_sheet") == r["tip_sheet_name"]:
                        entry = e
                        break
                if entry and entry["canonical"] in cc["excluded_canonicals"] and not r["excluded_from_pool"]:
                    r["excluded_from_pool"] = True
                    r["exclusion_reason"] = "manager_cut"
                    r["needs_review"] = True
                    r["review_reasons"].append(
                        f"Manager reported {entry['canonical']} as cut from the tip pool tonight -- "
                        "excluded from the Copy-for-Excel output."
                    )
        # Same cut cross-check for support staff (Busser/Runner/Barista/
        # Barback) -- previously only server_review got this treatment, so
        # a manager-reported cut for anyone in these four roles was
        # silently ignored and they stayed in the pastable output (real
        # case: a Runner, cut 7/22/26 but still showed
        # up -- see [[avra-mapping-issues]] memory for the write-up).
        if result.get("support_staff_review") is not None:
            for r in result["support_staff_review"]:
                entry = None
                for e in mapping.entries:
                    if e.get("tip_sheet") == r["employee"]:
                        entry = e
                        break
                if entry and entry["canonical"] in cc["excluded_canonicals"] and not r.get("cut_by_manager"):
                    r["cut_by_manager"] = True
                    r["needs_review"] = True
                    r["adp_flags"].append(
                        f"Manager reported {entry['canonical']} as cut from the tip pool tonight -- "
                        "excluded from the Copy-for-Excel output."
                    )

    # Legacy weekly-PDF-OCR cross-check path -- superseded by the daily
    # roster review above, kept only for the (currently unused) old
    # hotsched_pdf_path input so it doesn't need to be ripped out to add
    # the new flow. Only runs if that old param is actually passed in.
    if hotsched_pdf_path and not tip_sheet_path:
        result["partial_warning"] = "HotSchedules PDF uploaded, but the Tip Sheet wasn't -- the cross-check needs both. Go back and add the filled-in Tip Sheet."
    elif tip_sheet_path and hotsched_pdf_path:
        n_pages = ocr_hotschedules_pdf(hotsched_pdf_path, work_dir)
        hotsched_people = parse_all_hotsched_pages(work_dir, n_pages)
        tip_roles = load_tip_sheet_roles(tip_sheet_path)
        cross_report, unmatched_hs = run_hotsched_cross_check(tip_roles, hotsched_people, mapping)
        result["cross_check"] = cross_report
        result["cross_check_summary"] = dict(Counter(r["status"] for r in cross_report))
        result["unmatched_hotsched_names"] = sorted(unmatched_hs)
        result["hotsched_pages_processed"] = n_pages

    return result
