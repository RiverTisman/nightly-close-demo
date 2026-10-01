"""
Reads one day+shift tab from the real weekly Tip Sheet workbook (the
manager's own "WE M.DD.YY.xlsx" file, one workbook per week with a tab per
day+shift) to reconcile it against ADP.

Confirmed real layout, verified directly against WE 7.12.26.xlsx's MON AM
and MON PM tabs (identical structure on both): each tab has six fixed
Name-column ranges, one per role, matching the workbook's own Excel
dropdown-validation ranges (see EmployeeMapping.tip_sheet_entry /
build_support_staff_review in logic.py for why exact text matters here --
these are literal dropdown cells, not free text).

  Server:    A8:A43
  Bartender: A50:A65
  Runner:    A72:A86
  Barback:   J72:J86
  Busser:    A94:A124
  Barista:   J94:J124

Sheet names follow "{WEEKDAY} {SHIFT}", e.g. "MON PM" -- the weekday
abbreviations are the workbook's own (MON/TUES/WED/THURS/FRI/SAT/SUN), not
Python's %a.
"""

from datetime import date

import openpyxl

ROLE_RANGES = {
    "Server": ("A", 8, 43),
    "Bartender": ("A", 50, 65),
    "Runner": ("A", 72, 86),
    "Barback": ("J", 72, 86),
    "Busser": ("A", 94, 124),
    "Barista": ("J", 94, 124),
}

WEEKDAY_ABBR = {0: "MON", 1: "TUES", 2: "WED", 3: "THURS", 4: "FRI", 5: "SAT", 6: "SUN"}


class TipSheetTabNotFound(ValueError):
    pass


def sheet_name_for(date_obj, shift):
    return f"{WEEKDAY_ABBR[date_obj.weekday()]} {shift}"


def _sheet_date(cell_value):
    """Parses the sheet's own "M.D.YY" date string (cell A4, confirmed
    real format e.g. "7.6.26", "7.9.26") into a date, or None if it's
    missing/unparseable."""
    if not cell_value:
        return None
    try:
        month, day, year = (int(x) for x in str(cell_value).strip().split("."))
        return date(year + 2000 if year < 100 else year, month, day)
    except (ValueError, TypeError):
        return None


def read_weekly_totals_roster(xlsx_path):
    """Returns [(eid, name), ...] for every row in the WEEKLY TOTALS sheet
    that has both an EID (column A) and a name (column B) -- the exact
    source Excel's own Name-column dropdowns pull from, so it's ground
    truth for whether employee_mapping.json's tip_sheet text will register
    as "selected" when pasted, not just land as unselected typed text.

    Independent of which week's workbook this is or which date/shift is
    being checked -- WEEKLY TOTALS carries every currently active employee
    regardless of the specific week, unlike parse_tip_sheet_roster's
    per-day tabs.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "WEEKLY TOTALS" not in wb.sheetnames:
        raise TipSheetTabNotFound(
            "No \"WEEKLY TOTALS\" sheet in this workbook -- confirm this is a real Tip Sheet "
            "weekly workbook, not something else."
        )
    ws = wb["WEEKLY TOTALS"]
    roster = []
    for row in ws.iter_rows(min_col=1, max_col=2):
        eid_cell, name_cell = row[0], row[1]
        eid, name = eid_cell.value, name_cell.value
        if eid is None or not name or not str(name).strip():
            continue
        roster.append((str(eid), str(name)))
    return roster


def parse_tip_sheet_roster(xlsx_path, date_obj, shift):
    """Returns {role: [name strings, exactly as typed in that cell]} for
    the one day+shift tab matching date_obj/shift."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet_name = sheet_name_for(date_obj, shift)
    if sheet_name not in wb.sheetnames:
        raise TipSheetTabNotFound(
            f"No \"{sheet_name}\" tab in this workbook -- confirm this is the right week's "
            f"Tip Sheet and the date/shift match what you're checking. Tabs found: "
            f"{', '.join(wb.sheetnames)}"
        )
    ws = wb[sheet_name]

    # Tabs are named by weekday only (e.g. "THURS PM"), not by calendar
    # date -- every weekly workbook has one, so matching the tab name
    # alone can't catch "uploaded the wrong week's file." Cross-check the
    # sheet's own printed date (cell A4) against the date actually being
    # checked instead.
    sheet_dt = _sheet_date(ws["A4"].value)
    if sheet_dt and sheet_dt != date_obj.date():
        raise TipSheetTabNotFound(
            f"The \"{sheet_name}\" tab in this workbook is dated {ws['A4'].value} "
            f"({sheet_dt:%B %d, %Y}), not {date_obj:%B %d, %Y} -- this looks like the wrong "
            f"week's Tip Sheet. Upload the workbook that actually covers {date_obj:%B %d, %Y}."
        )
    roster = {}
    for role, (col, start, end) in ROLE_RANGES.items():
        names = []
        for r in range(start, end + 1):
            val = ws[f"{col}{r}"].value
            if val and str(val).strip():
                names.append(str(val))
        roster[role] = names
    return roster
